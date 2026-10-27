"""관리자 데이터 관리 공통 틀: 엑셀·CSV 내려받기 → 고쳐서 올리기 → 미리보기 → 반영 (2026-10-21 사용자 결정).

- 데이터(인적·교육)마다 시트 여러 개, 한글 열 이름. 엑셀 첫 시트 '안내'에 열 설명과 규칙을 적는다.
- 올리면 바로 바꾸지 않고 미리보기: 행마다 추가·변경·삭제·그대로·오류. 오류가 하나라도 있으면 반영할 수 없다.
- 파일에 없는 행은 그대로 둔다. '삭제' 열에 Y 를 적은 행만 지운다(직원은 퇴직 처리 · 쓰는 중인 자료는 목록에서만 뺀다).
- 파일에 없는 열은 바꾸지 않는다. 있는 열의 빈 칸은 '값 없음'으로 바꾼다(필수 열은 오류).
- 반영은 한 트랜잭션. 미리보기 뒤 데이터가 바뀌었으면 반영하지 않고 다시 올리게 한다.

데이터별 모듈(hr·education)은 SHEETS·current·check·apply 를 둔다.
"""
from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any, Callable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from ..errors import ApiError

DELETE = '삭제'
YES, NO = ('Y', 'YES', '예', 'O', 'TRUE', '1', '네'), ('N', 'NO', '아니오', 'X', 'FALSE', '0', '아니요')


@dataclass
class Col:
    header: str                 # 한글 열 이름(파일의 머리글)
    field: str                  # 내부 이름
    kind: str = 'text'          # text | int | num | date | bool | choice | list
    required: bool = False
    choices: tuple = ()         # choice: 화면에 보이는 값(내부 값과 같으면 그대로)
    help: str = ''
    readonly: bool = False      # 참고용(내려받기에만 의미 · 올려도 바꾸지 않음)
    width: int = 14
    example: str = ''
    default: Any = None         # 빈 칸일 때 값(예: 추천 사용은 Y)


@dataclass
class Sheet:
    name: str
    key: tuple                  # 행을 찾는 열(field) 이름
    cols: list
    help: str = ''
    auto_key: bool = False      # 열쇠를 비우면 새 행으로 보고 반영 때 번호를 정한다
    deletable: bool = True

    def col(self, f: str) -> Col:
        return next(c for c in self.cols if c.field == f)

    @property
    def editable(self) -> list:
        return [c for c in self.cols if not c.readonly and c.field not in self.key]


@dataclass
class Row:
    """올린 파일의 한 행(정리한 값)과 미리보기 결과"""
    sheet: str
    rowno: int
    values: dict
    present: set
    delete: bool = False
    key: tuple = ()
    action: str = 'same'        # add | update | delete | same | error | skip
    changes: list = field(default_factory=list)
    errors: list = field(default_factory=list)
    warnings: list = field(default_factory=list)
    tags: list = field(default_factory=list)
    label: str = ''

    def err(self, col: str, msg: str) -> None:
        self.errors.append({'col': col, 'msg': msg})
        self.action = 'error'

    def warn(self, msg: str) -> None:
        if msg not in self.warnings:
            self.warnings.append(msg)

    def view(self) -> dict:
        return {'sheet': self.sheet, 'row': self.rowno, 'key': ' · '.join(str(k) for k in self.key if k not in (None, '')) or '(새 행)',
                'label': self.label, 'action': self.action, 'changes': self.changes, 'errors': self.errors, 'warnings': self.warnings, 'tags': self.tags}

    def state(self) -> dict:
        """저장용(반영 때 같은 값으로 다시 계산한다)"""
        return {'sheet': self.sheet, 'rowno': self.rowno, 'values': _jsonable(self.values), 'present': sorted(self.present), 'delete': self.delete}


def _jsonable(v):
    if isinstance(v, dict):
        return {k: _jsonable(x) for k, x in v.items()}
    if isinstance(v, (list, tuple)):
        return [_jsonable(x) for x in v]
    if isinstance(v, (date, datetime)):
        return v.isoformat()[:10]
    return v


# ── 값 정리 ──
def show(c: Col, v) -> str:
    """화면·파일에 쓰는 글자"""
    if v is None or v == '' or v == []:
        return ''
    if c.kind == 'bool':
        return 'Y' if v else 'N'
    if c.kind == 'list':
        return '; '.join(str(x) for x in v)
    if c.kind == 'num':
        f = float(v)
        return str(int(f)) if f.is_integer() else f'{f:g}'
    if c.kind == 'date':
        return str(v)[:10]
    return str(v)


def parse(c: Col, raw) -> tuple[Any, str | None]:
    """파일 칸 → 값. (값, 오류 문구)"""
    if raw is None or (isinstance(raw, str) and not raw.strip()):
        return (c.default if c.default is not None else [] if c.kind == 'list' else None), None
    if isinstance(raw, float) and raw.is_integer() and c.kind in ('text', 'choice'):
        raw = int(raw)  # 엑셀이 숫자로 바꾼 사번·코드
    s = raw.strftime('%Y-%m-%d') if isinstance(raw, (datetime, date)) else str(raw).strip()
    if c.kind == 'int':
        try:
            f = float(s.replace(',', ''))
        except ValueError:
            return None, '숫자를 적어 주세요'
        if not f.is_integer():
            return None, '정수를 적어 주세요'
        return int(f), None
    if c.kind == 'num':
        try:
            return float(s.replace(',', '')), None
        except ValueError:
            return None, '숫자를 적어 주세요'
    if c.kind == 'date':
        m = re.fullmatch(r'(\d{4})[-./](\d{1,2})[-./](\d{1,2})\.?', s)
        if not m:
            return None, '날짜는 2024-03-02 처럼 적어 주세요'
        try:
            return date(int(m[1]), int(m[2]), int(m[3])).isoformat(), None
        except ValueError:
            return None, '없는 날짜예요'
    if c.kind == 'bool':
        u = s.upper()
        if u in YES:
            return True, None
        if u in NO:
            return False, None
        return None, 'Y 또는 N 으로 적어 주세요'
    if c.kind == 'choice':
        if c.choices and s not in c.choices:
            return None, f"다음 중 하나로 적어 주세요: {', '.join(c.choices)}"
        return s, None
    if c.kind == 'list':
        return [x.strip() for x in re.split(r'[;\n]', s) if x.strip()], None
    return s, None


def same(c: Col, a, b) -> bool:
    if c.kind == 'num':
        return (a is None and b is None) or (a is not None and b is not None and abs(float(a) - float(b)) < 1e-9)
    if c.kind == 'list':
        return list(a or []) == list(b or [])
    if c.kind == 'bool':
        return bool(a) == bool(b)
    return (a if a not in ('', None) else None) == (b if b not in ('', None) else None)


# ── 파일 쓰기 ──
HEAD_FILL = PatternFill('solid', fgColor='E8F1EE')
REQ_FILL = PatternFill('solid', fgColor='CFE6DD')
RO_FILL = PatternFill('solid', fgColor='EEEEEE')


def to_xlsx(title: str, guide: list[str], sheets: list[Sheet], rows: dict[str, list[dict]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = '안내'
    ws.column_dimensions['A'].width = 22
    ws.column_dimensions['B'].width = 10
    ws.column_dimensions['C'].width = 90
    ws.append([title])
    ws['A1'].font = Font(bold=True, size=14)
    for g in guide:
        ws.append([g])
    for sh in sheets:
        ws.append([])
        ws.append([f'■ {sh.name} 시트', '', sh.help])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
        ws.append(['열 이름', '구분', '설명'])
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)
            cell.fill = HEAD_FILL
        for c in sh.cols:
            kind = '열쇠' if c.field in sh.key else '참고용(바꿔도 반영 안 됨)' if c.readonly else '필수' if c.required else '선택'
            desc = c.help + (f" (선택지: {', '.join(c.choices)})" if c.choices else '') + (f' 예: {c.example}' if c.example else '')
            ws.append([c.header, kind, desc.strip()])
        if sh.deletable:
            ws.append([DELETE, '선택', '이 행을 지우려면 Y. 비워 두면 그대로예요.'])
    for r in ws.iter_rows(min_row=2):
        for cell in r:
            cell.alignment = Alignment(wrap_text=True, vertical='top')
    for sh in sheets:
        w = wb.create_sheet(sh.name)
        cols = sh.cols + ([Col(DELETE, '_delete', 'bool', width=8)] if sh.deletable else [])
        w.append([c.header for c in cols])
        for i, c in enumerate(cols, 1):
            cell = w.cell(1, i)
            cell.font = Font(bold=True)
            cell.fill = RO_FILL if c.readonly else REQ_FILL if (c.required or c.field in sh.key) else HEAD_FILL
            w.column_dimensions[get_column_letter(i)].width = c.width
            if c.kind in ('text', 'choice') or c.field in sh.key:
                w.column_dimensions[get_column_letter(i)].number_format = '@'
        for r in rows.get(sh.name, []):
            w.append([show(c, r.get(c.field)) if c.field != '_delete' else '' for c in cols])
        last = max(2, w.max_row + 200)
        for i, c in enumerate(cols, 1):
            opts = list(c.choices) if c.kind == 'choice' and c.choices and len(','.join(c.choices)) < 240 else ['Y', 'N'] if c.kind == 'bool' else None
            if opts and not c.readonly:
                dv = DataValidation(type='list', formula1='"' + ','.join(opts) + '"', allow_blank=True)
                w.add_data_validation(dv)
                dv.add(f'{get_column_letter(i)}2:{get_column_letter(i)}{last}')
        w.freeze_panes = 'A2'
        w.auto_filter.ref = f'A1:{get_column_letter(len(cols))}{max(1, w.max_row)}'
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def to_csv(sheet: Sheet, rows: list[dict]) -> bytes:
    cols = sheet.cols + ([Col(DELETE, '_delete', 'bool')] if sheet.deletable else [])
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator='\r\n')
    w.writerow([c.header for c in cols])
    for r in rows:
        w.writerow([show(c, r.get(c.field)) if c.field != '_delete' else '' for c in cols])
    return ('﻿' + buf.getvalue()).encode('utf-8')  # 엑셀이 한글을 바로 읽게 BOM


# ── 파일 읽기 ──
def _norm_head(h) -> str:
    return re.sub(r'[\s*]', '', str(h or ''))


def _table(sheet: Sheet, header: list, data: list[list], notes: list[str]) -> list[Row]:
    cols = sheet.cols + ([Col(DELETE, '_delete', 'bool')] if sheet.deletable else [])
    by_head = {_norm_head(c.header): c for c in cols}
    idx: dict[int, Col] = {}
    unknown = []
    for i, h in enumerate(header):
        c = by_head.get(_norm_head(h))
        if c:
            idx[i] = c
        elif str(h or '').strip():
            unknown.append(str(h).strip())
    if unknown:
        notes.append(f"{sheet.name} 시트: 모르는 열은 무시했어요 ({', '.join(unknown[:6])})")
    missing_keys = [sheet.col(k).header for k in sheet.key if sheet.col(k) not in idx.values()]
    if missing_keys and not sheet.auto_key:
        raise ApiError(422, 'MISSING_KEY_COLUMN', f"{sheet.name} 시트에 ‘{', '.join(missing_keys)}’ 열이 없어요", '내려받은 양식의 머리글(첫 줄)을 그대로 두고 고쳐 주세요')
    out = []
    for n, line in enumerate(data, 2):
        if not any(str(v).strip() for v in line if v is not None):
            continue
        r = Row(sheet.name, n, {}, set())
        for i, c in idx.items():
            raw = line[i] if i < len(line) else None
            if c.field == '_delete':
                v, e = parse(c, raw)
                r.delete = bool(v) if not e else False
                if e:
                    r.err(DELETE, e)
                continue
            if c.readonly:
                continue
            v, e = parse(c, raw)
            if e:
                r.err(c.header, e)
            r.values[c.field] = v
            r.present.add(c.field)
        r.key = tuple(r.values.get(k) for k in sheet.key)
        out.append(r)
    return out


def read_upload(name: str, raw: bytes, sheets: list[Sheet], sheet_hint: str | None) -> tuple[dict[str, list[Row]], list[str]]:
    """엑셀(여러 시트) 또는 CSV(시트 하나) → 시트별 행"""
    notes: list[str] = []
    found: dict[str, list[Row]] = {}
    low = name.lower()
    if low.endswith(('.xlsx', '.xlsm')):
        try:
            wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        except Exception as e:  # noqa: BLE001 - 깨진 파일은 사용자에게 알리고 끝낸다
            raise ApiError(422, 'BAD_XLSX', '엑셀 파일을 열지 못했어요', '내려받은 양식을 엑셀에서 저장한 .xlsx 파일인지 확인해 주세요', raw=str(e)[:200]) from e
        by_name = {_norm_head(s.name): s for s in sheets}
        for ws in wb.worksheets:
            sh = by_name.get(_norm_head(ws.title))
            if not sh:
                if ws.title != '안내':
                    notes.append(f"‘{ws.title}’ 시트는 이 데이터의 시트가 아니라 읽지 않았어요")
                continue
            it = ws.iter_rows(values_only=True)
            header = list(next(it, []) or [])
            found[sh.name] = _table(sh, header, [list(x) for x in it], notes)
        if not found:
            raise ApiError(422, 'NO_SHEET', f"읽을 시트가 없어요 (시트 이름: {', '.join(s.name for s in sheets)})", '내려받은 양식의 시트 이름을 바꾸지 말고 올려 주세요')
    elif low.endswith(('.csv', '.txt')):
        text = None
        for enc in ('utf-8-sig', 'cp949'):
            try:
                text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            raise ApiError(422, 'BAD_CSV', 'CSV 글자를 읽지 못했어요', 'UTF-8 또는 엑셀 기본(CP949) CSV로 저장해 주세요')
        lines = list(csv.reader(io.StringIO(text)))
        if not lines:
            raise ApiError(422, 'EMPTY', '빈 파일이에요')
        header = lines[0]
        heads = {_norm_head(h) for h in header}
        sh = next((s for s in sheets if s.name == sheet_hint), None) if sheet_hint else None
        if not sh:
            # 머리글이 가장 많이 맞는 시트
            score = lambda s: len(heads & {_norm_head(c.header) for c in s.cols})  # noqa: E731
            sh = max(sheets, key=score)
            if score(sh) < 2:
                raise ApiError(422, 'UNKNOWN_CSV', '어느 시트의 CSV인지 알 수 없어요', '시트를 고른 뒤 올리거나, 내려받은 CSV의 머리글을 그대로 두고 고쳐 주세요')
            notes.append(f'머리글로 ‘{sh.name}’ 시트의 CSV로 읽었어요')
        found[sh.name] = _table(sh, header, lines[1:], notes)
    else:
        raise ApiError(422, 'BAD_FORMAT', '엑셀(.xlsx) 또는 CSV(.csv) 파일만 올릴 수 있어요')
    return found, notes


# ── 미리보기 계산 ──
def diff(sheet: Sheet, rows: list[Row], current: dict[tuple, dict], label: Callable[[dict], str]) -> None:
    """행마다 추가·변경·삭제·그대로를 정한다. 같은 열쇠가 두 번 나오면 오류"""
    seen: dict[tuple, int] = {}
    for r in rows:
        blank_key = all(k in (None, '') for k in r.key)
        if blank_key and not sheet.auto_key:
            r.err(' · '.join(sheet.col(k).header for k in sheet.key), '열쇠 칸을 적어 주세요')
            continue
        if not blank_key:
            if r.key in seen:
                r.err(sheet.col(sheet.key[0]).header, f'{seen[r.key]}행과 열쇠가 같아요(한 번만 적어 주세요)')
                continue
            seen[r.key] = r.rowno
        cur = None if blank_key else current.get(r.key)
        r.label = label(dict(cur or {}, **{k: v for k, v in r.values.items() if v not in (None, '')}))
        if r.errors:
            continue
        if r.delete:
            if cur is None:
                r.action = 'skip'
                r.warn('지금 없는 행이라 지울 것이 없어요')
            else:
                r.action = 'delete' if sheet.deletable else 'error'
            continue
        if cur is None:
            r.action = 'add'
            for c in sheet.editable:
                if c.required and r.values.get(c.field) in (None, '', []):
                    r.err(c.header, '새 행에는 꼭 적어야 해요')
            r.changes = [{'col': c.header, 'before': '', 'after': show(c, r.values.get(c.field))} for c in sheet.editable
                         if c.field in r.present and r.values.get(c.field) not in (None, '', [], False)]
            continue
        changes = []
        for c in sheet.editable:
            if c.field not in r.present:
                continue
            new = r.values.get(c.field)
            if c.required and new in (None, '', []):
                r.err(c.header, '비울 수 없어요')
                continue
            if not same(c, cur.get(c.field), new):
                before, after = clip_pair(show(c, cur.get(c.field)), show(c, new))
                changes.append({'col': c.header, 'field': c.field, 'before': before, 'after': after})
        if r.errors:
            continue
        r.changes = changes
        r.action = 'update' if changes else 'same'


def clip_pair(a: str, b: str, room: int = 60) -> tuple[str, str]:
    """긴 글(규정 본문 등)은 바뀐 부분 앞뒤만 남긴다"""
    if len(a) <= 160 and len(b) <= 160:
        return a, b
    i = 0
    while i < min(len(a), len(b)) and a[i] == b[i]:
        i += 1
    j = 0
    while j < min(len(a), len(b)) - i and a[len(a) - 1 - j] == b[len(b) - 1 - j]:
        j += 1
    def cut(s: str) -> str:
        head = ('…' if i > room else '') + s[max(0, i - room):i]
        mid = s[i:len(s) - j]
        tail = s[len(s) - j:len(s) - j + room] + ('…' if j > room else '')
        return head + (mid if len(mid) <= 400 else mid[:400] + '…') + tail
    return cut(a), cut(b)


def merged(r: Row, cur: dict | None) -> dict:
    """반영할 최종 값: 지금 값 + 파일에 있는 열"""
    out = dict(cur or {})
    out.update({k: v for k, v in r.values.items() if k in r.present})
    return out


def summary(sheets: dict[str, list[Row]]) -> dict:
    out = {}
    for name, rows in sheets.items():
        cnt = {a: 0 for a in ('add', 'update', 'delete', 'same', 'error', 'skip')}
        for r in rows:
            cnt[r.action] = cnt.get(r.action, 0) + 1
        out[name] = cnt
    return out


def signature(sheets: dict[str, list[Row]]) -> list:
    """미리보기와 반영 때 계산이 같은지 비교하는 값"""
    return sorted([r.sheet, str(r.key), r.action, sorted((c['col'], c['after']) for c in r.changes)] for rs in sheets.values() for r in rs if r.action != 'same')


def rows_from_state(states: list[dict]) -> dict[str, list[Row]]:
    out: dict[str, list[Row]] = {}
    for s in states:
        out.setdefault(s['sheet'], []).append(Row(s['sheet'], s['rowno'], dict(s['values']), set(s['present']), s['delete']))
    return out
