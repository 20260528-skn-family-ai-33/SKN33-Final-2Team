"""5차 목업 DB 준비: 팀원 전체 백업 복원 → 보조 스키마(app_ext) → 보충 자료.

  python tools/setup_db.py            # 빈 Docker DB면 복원부터(백업 기본 위치 team_db/), 아니면 보충만 (여러 번 실행해도 같은 결과)
  python tools/setup_db.py --supplement-only   # 팀원이 운영하는 DB(복사본)에 붙을 때: 복원 없이 보조 스키마·보충만 · 그 DB 팀 표에 데이터를 쓴다

보충 자료 (2026-10-21 사용자 결정):
  - 관리자 전용 계정 1개(아이디 .env GP5_ADMIN_LOGIN · 비밀번호 GP5_DEMO_PASSWORD)
  - 팀 직원 가데이터 2.9.0 의 16명: 부서 계층·부서장·직급·인사 보조 항목·본인 입력·평정·학력·보유 자격 연결.
    직무는 팀 DB 직무표와 맞는 이름이 없어 job_id 를 채우지 않고 직무명만 글자로 둔다(팀 DB 원칙: 미확정 값은 추정하지 않음)
  - 팀 테스트 직원 14명의 비밀번호를 이 로컬 사본에서만 시연 비밀번호로 바꾼다
    (--keep-team-passwords 면 그대로 · --supplement-only 는 남의 DB라 늘 그대로 · 2026-10-23 사용자 결정)
  - 학습 자료 보조 값(형식·시간·단계·비용)을 규칙으로 채우고 '추정'으로 표시, 자격증 분야 묶음
  - 규정 묶음(적용 직원 구분·국외근무 전용·짧은 키)
  - 목표 후보 첫 생성(가데이터 규칙 · 본인 입력이 없으면 기본 후보)
"""
from __future__ import annotations

import argparse
import re
import subprocess
import sys
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

import openpyxl  # noqa: E402

from app import db  # noqa: E402
from app.auth import hash_password, verify_password  # noqa: E402
from app.config import S, env  # noqa: E402

# 팀 직원 가데이터 2.9.0: 이 폴더 data/seed 사본을 먼저, 없으면 팀원 DB 저장소 사본
XLSX_DEFAULT = next((p for p in (ROOT / 'data/seed/사내DB_직원가상데이터_구성초안_2.9.0.xlsx',
                                 ROOT.parent / 'project_full_20261021_151824/project/database/data/regulations/사내DB_직원가상데이터_구성초안_2.9.0.xlsx') if p.exists()),
                    ROOT / 'data/seed/사내DB_직원가상데이터_구성초안_2.9.0.xlsx')
# 복원 대상 컨테이너 = '{compose 이름}-db-1'. 고정값이면 다른 compose 프로젝트의 DB에 복원할 수 있어 설정으로 뺐다(2026-10-26)
COMPOSE_PROJECT = env('COMPOSE_PROJECT_NAME', 'gp5')
# 팀 DB 백업 기본 위치(.env TEAM_DB_DUMP 가 비면): 팀원이 받은 파일을 이 폴더 team_db/ 에 넣기만 하면 되게(2026-10-23 사용자 결정)
DUMP_DEFAULT = 'team_db/growpath_ai_it_20261022.dump'


def say(msg: str) -> None:
    print('  ' + msg, flush=True)


# ── 1. 복원 ──
def db_is_empty() -> bool:
    return db.val("SELECT count(*) FROM information_schema.tables WHERE table_schema = 'public' AND table_type = 'BASE TABLE'") == 0


def dump_path() -> Path:
    p = Path(env('TEAM_DB_DUMP', DUMP_DEFAULT))
    return p if p.is_absolute() else (ROOT / p).resolve()


def restore(dump: Path) -> None:
    if not dump.is_file():
        others = sorted(x.name for x in (ROOT / 'team_db').glob('*.dump'))
        raise SystemExit('\n'.join([
            f'팀 DB 백업 파일이 없어요: {dump}',
            f'  · 팀 DB 담당 팀원에게 백업 파일({dump.name})을 받아 위 위치에 넣고 다시 실행하세요. 기본 위치는 이 폴더의 team_db/ 예요.',
            '  · 다른 곳에 두었다면 .env 의 TEAM_DB_DUMP 에 그 위치를 적으세요.',
            *([f"  · team_db/ 에 있는 다른 백업: {', '.join(others)} — 이 파일을 쓰려면 TEAM_DB_DUMP 를 그 이름으로 고치세요(판이 다르면 시연 자료가 안 맞을 수 있어요)."] if others else []),
        ]))
    user, name = env('GP5_DB_USER', 'growpath'), env('GP5_DB_NAME', 'growpath')
    container = f'{COMPOSE_PROJECT}-db-1'
    say(f'팀원 백업 복원: {dump.name} → {container}')
    subprocess.run(['docker', 'cp', str(dump), f'{container}:/tmp/team.dump'], check=True)
    subprocess.run(['docker', 'exec', container, 'pg_restore', '-U', user, '-d', name, '--no-owner', '--no-privileges', '--exit-on-error', '/tmp/team.dump'], check=True)
    subprocess.run(['docker', 'exec', container, 'rm', '-f', '/tmp/team.dump'], check=True)


def apply_app_ext() -> None:
    sql = (ROOT / 'db/02_app_ext.sql').read_text(encoding='utf-8')
    # 매개변수 없이 보내야 SQL 안의 '%'(LIKE)를 자리표시자로 읽지 않는다
    raw = db.engine.raw_connection()
    try:
        raw.cursor().execute(sql)
        raw.commit()
    finally:
        raw.close()
    say('보조 스키마 app_ext 적용')


# ── 2. 계정 ──
def admin_account() -> None:
    with db.tx() as t:
        if not t.one('SELECT 1 FROM app_ext.admin_accounts WHERE login_id = :l', l=S.admin_login):
            t.run('INSERT INTO app_ext.admin_accounts(login_id, name, password_hash) VALUES(:l, :n, :h)', l=S.admin_login, n='시스템 관리자(시연)', h=hash_password(S.demo_password))
            say(f'관리자 계정 만듦: {S.admin_login}')


def team_passwords() -> None:
    n = 0
    with db.tx() as t:
        for e in t.all('SELECT employee_id, password_hash FROM public.employees WHERE employee_id NOT IN (SELECT employee_id FROM app_ext.employee_hr)'):
            if not verify_password(e['password_hash'], S.demo_password):
                t.run('UPDATE public.employees SET password_hash = :h WHERE employee_id = :i', h=hash_password(S.demo_password), i=e['employee_id'])
                n += 1
    say(f'팀 테스트 직원 비밀번호(로컬 사본만): {n}명 바꿈')


# ── 3. 직원 가데이터 2.9.0 ──
NONE_ROW = '(행 없음)'


def _lines(v) -> list[str]:
    s = '' if v is None else str(v).strip()
    return [] if s in ('', '0건', 'null', NONE_ROW, '없음') else [x.strip() for x in s.split('\n') if x.strip()]


def _text(v):
    s = None if v is None else str(v).strip()
    return None if s in (None, '', 'null', NONE_ROW, '미입력(열 생략)') else s


def _date(v):
    s = _text(v)
    return date.fromisoformat(s[:10]) if s and re.match(r'^\d{4}-\d{2}-\d{2}', s) else None


def _evals(v) -> list[dict]:
    """'기간 제목: 등급, 점수, 학습 사용 동의 예/아니오' (엑셀 열 이름은 옛 명칭 · 화면은 '목표 추천 활용 동의')"""
    out = []
    for ln in _lines(v):
        m = re.match(r'^(\S+)\s+(.*?):\s*(.*)$', ln)
        period, title, rest = (m.group(1), m.group(2), m.group(3)) if m else ('', ln, '')
        parts = [p.strip() for p in rest.split(',')]
        consent = next((p for p in parts if p.startswith('학습 사용 동의')), '')
        out.append({'period': period, 'title': title, 'grade': parts[0] if parts else '', 'score': parts[1] if len(parts) > 1 and parts[1] != '점수 없음' else '',
                    'consent': consent.endswith('예')})
    return out


LEVEL = {'학사': 'bachelor', '석사': 'master', '박사': 'doctorate', '전문학사': 'associate', '고졸': 'high_school'}
CERT_ALIAS = {'SQLD': 'SQL 개발자', 'SQLP': 'SQL 전문가', 'ADsP': '데이터 분석 준전문가', 'ADP': '데이터 분석 전문가', 'DAsP': '데이터아키텍처 준전문가',
              'DAP': '데이터아키텍처 전문가', '빅데이터분석기사': '빅데이터분석기사'}
GRADE_ORDER = {'1급': 1, '2급(갑)': 2, '2급(을)': 3, '3급': 4, '4급': 5, '5급': 6}


def employees(xlsx: Path) -> None:
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    rows = list(wb['직원'].iter_rows(values_only=True))
    head = [str(h) if h is not None else '' for h in rows[1]]
    col = lambda r, n: r[head.index(n)] if n in head else None  # noqa: E731
    people = [r for r in rows[5:] if r and r[0]]
    with db.tx() as t:
        inst = t.val("SELECT institution_id FROM public.institutions WHERE name = '한국수자원공사' LIMIT 1")
        # 부서: 엑셀 부서 키(BIZ·INFRA…)를 부서 코드로, 상위 부서는 이름으로 잇고 없으면 만든다(AI본부·경영부문)
        depts = {(str(col(r, 'department_id')), _text(col(r, 'department_label')), _text(col(r, '상위 부서'))) for r in people}
        code_of = {name: code for code, name, _ in depts}
        parent_of = {name: parent for _, name, parent in depts}
        tops = {p for p in parent_of.values() if p and p not in code_of}
        top_codes = {'AI본부': 'AIHQ', '경영부문': 'MGMT'}
        ids: dict[str, int] = {}

        def ensure(name: str, code: str, parent: str | None) -> int:
            if name in ids:
                return ids[name]
            pid = ensure(parent, code_of.get(parent) or top_codes.get(parent) or f'GP-{parent}', parent_of.get(parent)) if parent else None
            row = t.one('SELECT department_id FROM public.departments WHERE institution_id = :i AND department_code = :c', i=inst, c=code)
            if row:
                t.run('UPDATE public.departments SET name = :n, parent_department_id = :p WHERE department_id = :d', n=name, p=pid, d=row['department_id'])
                ids[name] = row['department_id']
            else:
                ids[name] = t.val('INSERT INTO public.departments(institution_id, department_code, name, parent_department_id) VALUES(:i, :c, :n, :p) RETURNING department_id',
                                  i=inst, c=code, n=name, p=pid)
            return ids[name]

        for name in sorted(tops):
            ensure(name, top_codes.get(name, f'GP-{name}'), None)
        for code, name, parent in sorted(depts):
            ensure(name, code, parent)
        # 직급: 팀 DB 이름 규칙('일반직 4급')을 따른다
        grade_ids = {}
        for g in sorted({_text(col(r, 'grade_label')) for r in people}, key=lambda x: GRADE_ORDER.get(x, 9)):
            gname = f'일반직 {g}'
            row = t.one('SELECT job_grade_id FROM public.job_grades WHERE institution_id = :i AND name = :n', i=inst, n=gname)
            grade_ids[g] = row['job_grade_id'] if row else t.val('INSERT INTO public.job_grades(institution_id, name, sort_order) VALUES(:i, :n, :s) RETURNING job_grade_id',
                                                                  i=inst, n=gname, s=GRADE_ORDER.get(g, 9))
        majors = {r['title']: r['major_id'] for r in t.all('SELECT major_id, title FROM public.majors')}
        certs = {r['title']: r['cert_id'] for r in t.all('SELECT cert_id, title FROM public.certifications')}
        pw = hash_password(S.demo_password)
        added = 0
        for r in people:
            num = str(col(r, 'employee_code'))
            dept = _text(col(r, 'department_label'))
            head_ = str(col(r, '상급자 권한') or '') == '예'
            row = t.one('SELECT employee_id FROM public.employees WHERE institution_id = :i AND employee_number = :n', i=inst, n=num)
            vals = dict(n=str(col(r, 'display_name')), em=f'{num}@kwater-virtual.example.invalid', jt=_text(col(r, '직무명')), i=inst, num=num,
                        d=ids[dept], g=grade_ids[_text(col(r, 'grade_label'))], h=head_)
            if row:
                eid = row['employee_id']
                t.run("""UPDATE public.employees SET name = :n, email = :em, job_title = :jt, department = NULL, grade = NULL, department_id = :d, job_grade_id = :g,
                         is_department_head = :h, updated_at = now() WHERE employee_id = :eid""", eid=eid, **{k: v for k, v in vals.items() if k not in ('i', 'num')})
            else:
                eid = t.val("""INSERT INTO public.employees(name, email, job_title, institution_id, employee_number, department_id, job_grade_id, is_department_head,
                               source_type, password_hash) VALUES(:n, :em, :jt, :i, :num, :d, :g, :h, 'TEST_MANUAL', :pw) RETURNING employee_id""", pw=pw, **vals)
                added += 1
            t.run("""INSERT INTO app_ext.employee_hr(employee_id, employee_group, hire_date, position, job_family, task, employment_status, account_status,
                         grade_since, valid_from, change_reason) VALUES(:e, :g, :hd, :p, :jf, :tk, :es, :ac, :gs, :vf, :cr)
                     ON CONFLICT (employee_id) DO UPDATE SET employee_group = EXCLUDED.employee_group, hire_date = EXCLUDED.hire_date, position = EXCLUDED.position,
                         job_family = EXCLUDED.job_family, task = EXCLUDED.task, employment_status = EXCLUDED.employment_status, account_status = EXCLUDED.account_status,
                         grade_since = EXCLUDED.grade_since, valid_from = EXCLUDED.valid_from, change_reason = EXCLUDED.change_reason, updated_at = now()""",
                  e=eid, g=_text(col(r, 'employee_group')), hd=_date(col(r, 'hire_date')), p=_text(col(r, 'position_label')), jf=_text(col(r, 'job_family_label')),
                  tk=_text(col(r, 'task_summary')), es=_text(col(r, 'employment_status')) or '재직', ac=_text(col(r, 'account_status')) or 'active',
                  gs=_date(col(r, '현 직급 임용일')), vf=_date(col(r, 'valid_from')), cr=_text(col(r, 'change_reason')))
            # 본인 입력(사내 DB 초기값). 직원이 내 정보에서 고친 값(edited)은 건드리지 않는다
            has = _text(col(r, 'work_description')) is not None or str(col(r, 'experience_status') or '') not in ('', NONE_ROW)
            st = lambda k: str(col(r, f'input_status.{k}') or 'unprovided')  # noqa: E731
            yrs = _text(col(r, 'development_years'))
            base = {'workDesc': _text(col(r, 'work_description')), 'expStatus': _text(col(r, 'experience_status')), 'expSummary': _text(col(r, 'experience_summary')),
                    'devYears': float(yrs) if yrs and re.match(r'^[\d.]+$', yrs) else None,
                    'inputStatus': {'certs': st('owned_certifications'), 'learning': st('learning_history'), 'projects': st('project_history'), 'tech': st('technology_stack')},
                    'tech': _lines(col(r, '기술 경험')), 'certs': _lines(col(r, '보유 자격')), 'learning': _lines(col(r, '교육 이력')), 'projects': _lines(col(r, '프로젝트')),
                    'present': has}
            t.run('INSERT INTO app_ext.employee_self_inputs(employee_id, baseline) VALUES(:e, :b) ON CONFLICT (employee_id) DO UPDATE SET baseline = EXCLUDED.baseline',
                  e=eid, b=db.J(base))
            for ev in _evals(col(r, '본인 입력 평가')):
                t.run("""INSERT INTO app_ext.employee_evaluations(employee_id, period, title, grade, score, consent_default) VALUES(:e, :p, :t, :g, :s, :c)
                         ON CONFLICT (employee_id, period) DO UPDATE SET title = EXCLUDED.title, grade = EXCLUDED.grade, score = EXCLUDED.score, consent_default = EXCLUDED.consent_default""",
                      e=eid, p=ev['period'], t=ev['title'], g=ev['grade'], s=ev['score'], c=ev['consent'])
            # 학력(ERD employee_education_records) — 학과 이름이 팀 DB 전공과 똑같을 때만 전공에 잇는다
            t.run('DELETE FROM public.employee_education_records WHERE employee_id = :e', e=eid)
            for ln in _lines(col(r, '학력')):
                m = re.match(r'^(\S+)\s+(\S+)\s*\((\S+?),', ln)
                if m:
                    t.run("""INSERT INTO public.employee_education_records(employee_id, school_name, education_level, department_name, major_id)
                             VALUES(:e, :s, :l, :d, :m)""", e=eid, s=m.group(1), l=LEVEL.get(m.group(3), 'other'), d=m.group(2), m=majors.get(m.group(2)))
            # 보유 자격(ERD employee_certifications) — 팀 DB 자격 이름과 맞는 것만
            t.run('DELETE FROM public.employee_certifications WHERE employee_id = :e', e=eid)
            for ln in base['certs']:
                key = ln.split(' (')[0].strip()
                cid = certs.get(CERT_ALIAS.get(key, key))
                if cid:
                    t.run('INSERT INTO public.employee_certifications(employee_id, certification_id) VALUES(:e, :c) ON CONFLICT DO NOTHING', e=eid, c=cid)
    say(f'직원 가데이터 {xlsx.name}: {len(people)}명 (새로 넣음 {added}명) · 부서 {len(ids)}개')


# ── 4. 학습 자료 보조 값 ──
DIFF_STAGE = {'BEGINNER': 1, 'INTERMEDIATE': 2, 'ADVANCED': 3}
AI_SKILLS = {'RAG & 생성형 AI 에이전트', '자연어 처리(NLP) & LLM 기초', '머신러닝 지도/비지도 학습', 'MLOps & 모델 서빙', 'ETL & 데이터 파이프라인',
             '탐색적 데이터 분석(EDA) & 전처리', '데이터 웨어하우스 & OLAP', '딥러닝 신경망 아키텍처', 'PyTorch/TensorFlow 딥러닝 구현', '컴퓨터 비전(CV) 모델링'}
CLOUD_SKILLS = {'클라우드 컴퓨팅 아키텍처 (AWS/Azure)', '클라우드 IAM & 네트워크 보안', '네트워크 보안 & 침입 탐지', '시스템 & OS 보안 강화',
                '데이터 거버넌스 & 보안 규정 준수', '웹 애플리케이션 취약점 진단', 'TCP/IP 네트워크 프로토콜'}


def resources() -> None:
    with db.tx() as t:
        n = 0
        for c in t.all('SELECT course_id, platform, cost, duration_weeks, difficulty FROM public.courses'):
            kmooc = (c['platform'] or '') == 'K-MOOC'
            hours = (c['duration_weeks'] * 2 if c['duration_weeks'] else 10) if kmooc else 1
            est = ['format', 'hours', 'stage'] + ([] if c['cost'] is not None else [])
            n += t.run("""INSERT INTO app_ext.resource_meta(resource_key, resource_type, course_id, format, hours, stage, cost_won, estimated_fields, is_real)
                          VALUES(:k, 'course', :cid, :f, :h, :s, :c, :e, true) ON CONFLICT (resource_key) DO NOTHING""",
                       k=c['course_id'], cid=c['course_id'], f='영상' if kmooc else '문서', h=hours, s=DIFF_STAGE.get(c['difficulty'] or '', 2),
                       c=int(c['cost']) if c['cost'] is not None else None, e=db.J(est))
        for b in t.all('SELECT isbn FROM public.books'):
            n += t.run("""INSERT INTO app_ext.resource_meta(resource_key, resource_type, isbn, format, hours, stage, estimated_fields, is_real)
                          VALUES(:k, 'book', :i, '문서', 10, 2, :e, true) ON CONFLICT (resource_key) DO NOTHING""", k=f"book_{b['isbn']}", i=b['isbn'], e=db.J(['format', 'hours', 'stage']))
        # 자격증 분야: 시행 기관 + 스킬(마이크로소프트는 스킬로 나눔)
        groups = [('kdata', '데이터·SQL (한국데이터산업진흥원)', '한국데이터산업진흥원', 'https://www.dataq.or.kr/www/sub/introduction.do', 1),
                  ('ms-ai', 'Microsoft · AI·데이터', 'Microsoft', 'https://learn.microsoft.com/credentials/', 2),
                  ('ms-cloud', 'Microsoft · 클라우드·보안', 'Microsoft', 'https://learn.microsoft.com/credentials/', 3),
                  ('ms-dev', 'Microsoft · 개발·운영', 'Microsoft', 'https://learn.microsoft.com/credentials/', 4),
                  ('ms-etc', 'Microsoft · 업무 앱·기타', 'Microsoft', 'https://learn.microsoft.com/credentials/', 5)]
        for g in groups:
            t.run("""INSERT INTO app_ext.cert_groups(group_key, label, issuer, url, display_order) VALUES(:k, :l, :i, :u, :o)
                     ON CONFLICT (group_key) DO UPDATE SET label = EXCLUDED.label, issuer = EXCLUDED.issuer, url = EXCLUDED.url, display_order = EXCLUDED.display_order""",
                  k=g[0], l=g[1], i=g[2], u=g[3], o=g[4])
        skills = {}
        for r in t.all('SELECT cs.cert_id, s.name FROM public.certification_skills cs JOIN public.skill_tags s USING (skill_id)'):
            skills.setdefault(r['cert_id'], set()).add(r['name'])
        for x in t.all('SELECT cert_id, title, provider FROM public.certifications ORDER BY cert_id'):
            sk = skills.get(x['cert_id'], set())
            if (x['provider'] or '') == '한국데이터산업진흥원':
                grp = 'kdata'
            elif (x['provider'] or '').startswith('Microsoft'):
                grp = 'ms-ai' if sk & AI_SKILLS else 'ms-cloud' if sk & CLOUD_SKILLS else 'ms-dev' if sk else 'ms-etc'
            else:
                key = 'org-' + re.sub(r'\W+', '-', x['provider'] or '기타').strip('-').lower()
                t.run('INSERT INTO app_ext.cert_groups(group_key, label, issuer, display_order) VALUES(:k, :l, :i, 9) ON CONFLICT DO NOTHING', k=key, l=x['provider'] or '기타', i=x['provider'])
                grp = key
            n += t.run("""INSERT INTO app_ext.resource_meta(resource_key, resource_type, cert_id, format, hours, stage, estimated_fields, is_real, cert_code, cert_group, cert_short)
                          VALUES(:k, 'cert', :c, '시험', 30, 4, :e, true, :code, :g, :short) ON CONFLICT (resource_key) DO NOTHING""",
                       k=f"cert_{x['cert_id']}", c=x['cert_id'], e=db.J(['hours']), code=str(x['cert_id']), g=grp, short=x['title'])
    say(f'학습 자료 보조 값(추정) 새로 넣음: {n}건')


# ── 5. 규정 묶음 ──
GENERAL = ['인재개발규정', '취업규칙', '유연근무제 운영기준', '근무및부업관리규정', '여비규정', '인사규정', '인사규정시행세칙', '연봉규정', '연봉규정시행세칙',
           '복지후생규정', '복지후생규정시행세칙', '직제규정', '직제규정 시행세칙', '위임전결규정', '인사규정 채용 세칙', '내부임용시험시행지침']
OVERSEAS = ['국외근무직원복무규정', '국외근무직원연봉및복지후생규정']
OTHER = {'실무직': ['실무직 관리규정', '실무직 관리기준', '실무직 보수기준'], '운영직': ['운영직관리세칙'], '전문직': ['전문직관리세칙'], '특정직': ['특정직관리세칙'],
         '청원경찰': ['청원경찰관리규정'], '기간제근로자': ['기간제근로자 관리규정', '기간제근로자 관리기준'], '특수직': ['특수직 보수기준']}
# 4차 조항 id 앞머리와 같은 짧은 키(화면 주소·기록에 씀)
SHORT = {'인재개발규정': 'dev', '취업규칙': 'emp', '유연근무제 운영기준': 'flex', '근무및부업관리규정': 'side', '여비규정': 'trip', '인사규정': 'hr',
         '인사규정시행세칙': 'hrd', '연봉규정': 'pay', '연봉규정시행세칙': 'payd', '복지후생규정': 'wel', '복지후생규정시행세칙': 'weld', '직제규정': 'org',
         '직제규정 시행세칙': 'orgd', '위임전결규정': 'dele', '국외근무직원복무규정': 'osw', '국외근무직원연봉및복지후생규정': 'ospay', '실무직 관리규정': 'prac',
         '실무직 관리기준': 'pracs', '실무직 보수기준': 'pracp', '인사규정 채용 세칙': 'hire', '내부임용시험시행지침': 'exam'}


def regulations() -> None:
    group_of = {n: '일반직' for n in GENERAL + OVERSEAS}
    for g, names in OTHER.items():
        for n in names:
            group_of[n] = g
    with db.tx() as t:
        n = 0
        for r in t.all('SELECT regulation_id, name FROM regulations.regulations ORDER BY name'):
            name = r['name']
            if t.one('SELECT 1 FROM app_ext.regulation_meta WHERE regulation_id = :i', i=r['regulation_id']):
                continue
            key = SHORT.get(name) or 'reg-' + r['regulation_id'][:8]
            t.run("""INSERT INTO app_ext.regulation_meta(regulation_id, short_key, employee_group, overseas_only, org) VALUES(:i, :k, :g, :o, '한국수자원공사')""",
                  i=r['regulation_id'], k=key, g=group_of.get(name, '미분류'), o=name in OVERSEAS)
            n += 1
    say(f'규정 묶음 정보 새로 넣음: {n}건')


def goals() -> None:
    try:
        from app.services import goals as G
    except ImportError as e:  # 서버 코드가 아직 없을 때(처음 개발 중)는 건너뛴다
        say(f'목표 후보 첫 생성 건너뜀: {e}')
        return
    n = G.seed_all()
    say(f'목표 후보 첫 생성(가데이터 규칙·기본 후보): {n}명')


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument('--supplement-only', action='store_true')
    ap.add_argument('--keep-team-passwords', action='store_true')
    ap.add_argument('--xlsx', default=str(XLSX_DEFAULT))
    args = ap.parse_args()
    print('GrowPath 5차 DB 준비')
    if not args.supplement_only:
        if db_is_empty():
            restore(dump_path())
        else:
            say('DB에 이미 표가 있어 복원은 건너뜀')
    apply_app_ext()
    # 시연 비밀번호·직원 가데이터는 저장소 밖에서 받는다(팀 database/RESTORE.md · 2026-10-26). 없으면 그 단계만 건너뛴다
    if not S.demo_password:
        say('GP5_DEMO_PASSWORD 가 비어 있어 관리자 계정·직원 가데이터·팀 테스트 비밀번호 단계를 건너뜀(.env 에 넣고 다시 실행)')
    else:
        admin_account()
        if Path(args.xlsx).exists():
            employees(Path(args.xlsx))
        else:
            say(f'직원 가데이터 엑셀이 없어 건너뜀: {args.xlsx} (담당자에게 받아 data/seed/ 에 넣고 다시 실행)')
        if args.supplement_only:
            say('팀 테스트 직원 비밀번호: 그대로 둠(--supplement-only · 팀원이 운영하는 DB)')
        elif not args.keep_team_passwords:
            team_passwords()
    resources()
    regulations()
    goals()
    print('완료')


if __name__ == '__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    main()
